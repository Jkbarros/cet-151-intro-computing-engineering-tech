#!/usr/bin/env python3
"""Convert selected Unit DOCX submissions to readable markdown in docs/."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

# Canonical copies (avoid duplicate Unit 1 folders in the export)
SOURCES: dict[str, str] = {
    "Unit 1 JBarros/Unit 1/Project 1A.docx": "unit-01-project-1a-invitation.md",
    "Unit 1 JBarros/Unit 1/Project 1B.docx": "unit-01-project-1b-header.md",
    "Unit 1 JBarros/Unit 1/Project 1C.docx": "unit-01-project-1c-resume.md",
    "Unit 2/Unit 2 JB/Project 2C.docx": "unit-02-project-2c-unit-conversions.md",
    "Test 1/Test 1/Test 1.docx": "test-1-notes.md",
}


def para_to_md(text: str) -> str:
    return text.strip()


def table_to_md(table) -> list[str]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
    if not rows:
        return []
    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join("---" for _ in rows[0]) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return lines


def convert_docx(src: Path, dest: Path, title: str) -> None:
    doc = Document(src)
    rel = src.relative_to(ROOT).as_posix()
    lines: list[str] = [
        f"# {title}",
        "",
        f"**Source:** [`{src.name}`](../{rel.replace(' ', '%20')})  ",
        "**Course:** CET 151 - Intro to Computing in Engineering Technology (EMU)  ",
        "**Author:** Julio Barros",
        "",
    ]
    for para in doc.paragraphs:
        t = para_to_md(para.text)
        if not t:
            lines.append("")
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style:
            lines.extend(["", f"## {t}", ""])
        elif "heading 2" in style:
            lines.extend(["", f"### {t}", ""])
        elif "list" in style:
            lines.append(f"- {t}")
        else:
            lines.append(t)
            lines.append("")

    for table in doc.tables:
        lines.extend(["", *table_to_md(table), ""])

    body = "\n".join(lines)
    body = re.sub(r"\n{3,}", "\n\n", body).strip() + "\n"
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(body, encoding="utf-8")


def main() -> None:
    for rel, slug in SOURCES.items():
        src = ROOT / rel
        if not src.exists():
            print(f"skip missing {rel}")
            continue
        title = src.stem.replace("Project ", "Project ").replace("Test ", "Test ")
        convert_docx(src, DOCS / slug, title)
        print(f"wrote docs/{slug}")


if __name__ == "__main__":
    main()
